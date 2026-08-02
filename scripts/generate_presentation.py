"""
Generate Jagiri's Kutumbam overview deck (PPTX + PDF).
Run: python scripts/generate_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "presentations"
PPTX_PATH = OUT_DIR / "Jagiris-Kutumbam-Overview.pptx"
PDF_PATH = OUT_DIR / "Jagiris-Kutumbam-Overview.pdf"

# ── Brand palette ─────────────────────────────────────────────────────────────
NAVY = RGBColor(0x0F, 0x17, 0x2A)
PRIMARY = RGBColor(0x25, 0x63, 0xEB)
PRIMARY_LIGHT = RGBColor(0xDB, 0xEA, 0xFE)
TEAL = RGBColor(0x08, 0x91, 0xB2)
TEAL_LIGHT = RGBColor(0xCC, 0xFB, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x64, 0x74, 0x8B)
SLATE_LIGHT = RGBColor(0x94, 0xA3, 0xB8)
DARK = RGBColor(0x1E, 0x29, 0x3B)
CANVAS = RGBColor(0xF1, 0xF5, 0xF9)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
EMERALD = RGBColor(0x05, 0x96, 0x69)
AMBER = RGBColor(0xD9, 0x77, 0x06)

PDF_NAVY = colors.HexColor("#0F172A")
PDF_PRIMARY = colors.HexColor("#2563EB")
PDF_TEAL = colors.HexColor("#0891B2")
PDF_SLATE = colors.HexColor("#64748B")
PDF_LIGHT = colors.HexColor("#F1F5F9")
PDF_CARD = colors.HexColor("#FFFFFF")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Segoe UI"

# Layout grid (inches)
ML = 0.72
MR = 0.72
HEADER_H = 1.38
FOOTER_H = 0.48
CONTENT_W = 13.333 - ML - MR
ROW_H = 0.74
DOT_COL = 0.42
TEXT_INDENT = ML + DOT_COL + 0.14

SLIDES: list[dict] = [
    {
        "kind": "cover",
        "title": "Jagiri's Kutumbam",
        "subtitle": "Family Memory Management Application",
        "tagline": "Connect · Remember · Celebrate · Grow Together",
        "footer": "Web & Android  ·  Private family portal",
    },
    {
        "kind": "agenda",
        "title": "Agenda",
        "items": [
            ("01", "Purpose & vision", PRIMARY),
            ("02", "Platform access", TEAL),
            ("03", "Core modules", INDIGO),
            ("04", "Admin & security", EMERALD),
            ("05", "Summary", AMBER),
        ],
    },
    {
        "kind": "bullets",
        "title": "What Is Jagiri's Kutumbam?",
        "intro": "A secure digital home for the Jagiri family to preserve identity, history, and connection.",
        "bullets": [
            "Central directory of family members with rich profiles",
            "Visual family tree built from parent–child relationships",
            "Photos, events, maps, and analytics in one place",
            "Role-based access — only approved members can sign in",
        ],
    },
    {
        "kind": "two_col",
        "title": "Platforms & Access",
        "left_title": "Where it runs",
        "left": [
            "Web app — modern browser (desktop & tablet)",
            "Android app — native shell via Capacitor",
            "Responsive layout with dark & light themes",
        ],
        "right_title": "Getting started",
        "right": [
            "Register or receive credentials from an admin",
            "Sign in with username & password",
            "Forgot password & privacy policy supported",
            "Notifications via the bell icon in the header",
        ],
    },
    {
        "kind": "module",
        "badge": "01",
        "badge_color": PRIMARY,
        "title": "Dashboard",
        "intro": "Your family command centre — key metrics and charts at a glance.",
        "bullets": [
            "Live counts: members, living/deceased, upcoming birthdays",
            "Charts: births by decade, gender, blood groups, deaths by year",
            "Recent announcements and activity feed",
            "Quick navigation to members, gallery, and events",
        ],
    },
    {
        "kind": "module",
        "badge": "02",
        "badge_color": TEAL,
        "title": "Family Members",
        "intro": "Complete, searchable directory of every family member.",
        "bullets": [
            "Add & edit profiles — photo, dates, parents, spouse, occupation",
            "Smart filters: village/place, surname, gender, living status",
            "Filters hidden by default — tap Filters to expand on mobile",
            "Tap any member to open a detailed profile page",
        ],
    },
    {
        "kind": "module",
        "badge": "03",
        "badge_color": INDIGO,
        "title": "Family Tree",
        "intro": "Interactive visual genealogy — automatically built from member links.",
        "bullets": [
            "Vertical tree with spouse pairs shown side-by-side",
            "Click any card to open the member profile",
            "Pan, scroll-to-zoom, Fit, Zoom in/out, and Full screen",
            "Optimized for both web and mobile viewports",
        ],
    },
    {
        "kind": "module",
        "badge": "04",
        "badge_color": EMERALD,
        "title": "Reports & Analytics",
        "intro": "Demographic insights with drill-down to member lists.",
        "bullets": [
            "Core: total, living, deceased, male, female members",
            "Age groups for living members (children to seniors)",
            "Additional: blood group, occupation, village, and more",
            "Tap any report card → view filtered member table",
        ],
    },
    {
        "kind": "module",
        "badge": "05",
        "badge_color": AMBER,
        "title": "Photo Gallery",
        "bullets": [
            "Shared family photo album accessible to all members",
            "Browse uploaded memories in a clean gallery grid",
            "Upload new photos from the gallery module",
            "Photos served securely through the family backend",
        ],
    },
    {
        "kind": "module",
        "badge": "06",
        "badge_color": PRIMARY,
        "title": "Events",
        "bullets": [
            "Family events calendar — gatherings, milestones, reminders",
            "Events surface on the dashboard upcoming list",
            "Helps the family stay aligned on important dates",
            "Easy to scan on mobile and desktop",
        ],
    },
    {
        "kind": "module",
        "badge": "07",
        "badge_color": TEAL,
        "title": "Places & Map",
        "intro": "See where family members live and explore locations visually.",
        "bullets": [
            "Interactive map with pins grouped by place/village",
            "Tap a pin → member list with profile links",
            "Open Google Maps directions from any location",
            "Filter panel for place and surname on mobile & web",
        ],
    },
    {
        "kind": "bullets",
        "title": "Search · Contact · Account",
        "bullets": [
            "Global Search — find members quickly across the app",
            "Contact Us — reach family organisers",
            "Account & Privacy — profile settings and data controls",
            "Privacy Policy page for transparency",
        ],
    },
    {
        "kind": "module",
        "badge": "08",
        "badge_color": EMERALD,
        "title": "Admin Portal",
        "intro": "Available to administrators only.",
        "bullets": [
            "User management — create, edit, and approve access",
            "Announcements — broadcast news to all members",
            "Audit log — track sensitive changes",
            "Keeps the family network secure and well-governed",
        ],
    },
    {
        "kind": "two_col",
        "title": "Technology Stack",
        "left_title": "Frontend & mobile",
        "left": [
            "React 18 + Vite",
            "Tailwind CSS, Recharts, Leaflet",
            "Capacitor 8 — Android app",
        ],
        "right_title": "Backend & data",
        "right": [
            "Node.js + Express REST API",
            "PostgreSQL (my-family database)",
            "JWT authentication & role-based routes",
        ],
    },
    {
        "kind": "closing",
        "title": "One App. One Family. All Connected.",
        "bullets": [
            "Members  ·  Tree  ·  Reports  ·  Gallery",
            "Events  ·  Map  ·  Search  ·  Notifications",
        ],
        "footer": "Thank you  ·  Questions welcome",
    },
]


# ── Low-level helpers ─────────────────────────────────────────────────────────
def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left, top, width, height, fill: RGBColor, *, radius=False, line=None):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _text(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _run(p, text, *, size=18, bold=False, color=DARK, align=None):
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return r


def _in(v: float):
    return Inches(v)


# ── Slide builder ─────────────────────────────────────────────────────────────
class SlideBuilder:
    def __init__(self, slide, page: int, total: int):
        self.slide = slide
        self.page = page
        self.total = total

    def canvas(self):
        _set_bg(self.slide, CANVAS)
        # subtle top accent strip
        _rect(self.slide, _in(0), _in(0), SLIDE_W, _in(0.06), PRIMARY)

    def footer(self):
        y = SLIDE_H - _in(FOOTER_H)
        _rect(self.slide, _in(0), y, SLIDE_W, _in(FOOTER_H), NAVY)
        left = _text(self.slide, _in(ML), y + _in(0.1), _in(5), _in(FOOTER_H - 0.2))
        ltf = left.text_frame
        ltf.clear()
        ltf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _run(ltf.paragraphs[0], "Jagiri's Kutumbam", size=11, color=SLATE_LIGHT)
        right = _text(self.slide, SLIDE_W - _in(MR + 1.2), y + _in(0.1), _in(1.2), _in(FOOTER_H - 0.2))
        rtf = right.text_frame
        rtf.clear()
        rtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rp = rtf.paragraphs[0]
        rp.alignment = PP_ALIGN.RIGHT
        _run(rp, f"{self.page} / {self.total}", size=11, bold=True, color=WHITE)

    def header(self, title: str, intro: str | None = None, badge: str | None = None, badge_color: RGBColor = PRIMARY):
        _rect(self.slide, _in(0), _in(0.06), SLIDE_W, _in(HEADER_H - 0.06), CARD)
        _rect(self.slide, _in(0), _in(HEADER_H), SLIDE_W, _in(0.04), PRIMARY)

        x_title = _in(ML)
        if badge:
            bx = _in(ML)
            by = _in(0.32)
            badge_shape = _rect(self.slide, bx, by, _in(0.62), _in(0.62), badge_color, radius=True)
            btf = badge_shape.text_frame
            btf.clear()
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            _run(bp, badge, size=14, bold=True, color=WHITE)
            x_title = _in(ML + 0.82)

        tb = _text(self.slide, x_title, _in(0.28), _in(CONTENT_W - (0.82 if badge else 0)), _in(0.55))
        tp = tb.text_frame.paragraphs[0]
        _run(tp, title, size=30, bold=True, color=NAVY)

        if intro:
            ib = _text(self.slide, _in(ML), _in(0.88), _in(CONTENT_W), _in(0.42))
            ip = ib.text_frame.paragraphs[0]
            _run(ip, intro, size=14, color=SLATE)

    def content_card(self, top: float, height: float):
        return _rect(
            self.slide,
            _in(ML),
            _in(top),
            _in(CONTENT_W),
            _in(height),
            CARD,
            radius=True,
            line=BORDER,
        )

    def bullet_rows(self, bullets: list[str], top: float = 1.58, card_top: float | None = None):
        """Grid-aligned bullet list — dot and text share fixed row heights."""
        n = len(bullets)
        card_t = card_top if card_top is not None else top - 0.12
        card_h = n * ROW_H + 0.28
        self.content_card(card_t, card_h)

        for i, text in enumerate(bullets):
            row_y = top + i * ROW_H
            row_mid = row_y + ROW_H / 2

            # connector line between dots (except last)
            if i < n - 1:
                line_x = _in(ML + 0.33)
                line_top = _in(row_y + ROW_H * 0.55)
                _rect(self.slide, line_x, line_top, _in(0.03), _in(ROW_H * 0.45), BORDER)

            # dot
            dot_size = 0.18
            dot = _rect(
                self.slide,
                _in(ML + 0.24),
                _in(row_mid - dot_size / 2),
                _in(dot_size),
                _in(dot_size),
                PRIMARY if i % 2 == 0 else TEAL,
                radius=True,
            )

            # text — vertically centred in row
            tb = _text(
                self.slide,
                _in(TEXT_INDENT),
                _in(row_y + 0.1),
                _in(CONTENT_W - DOT_COL - 0.36),
                _in(ROW_H - 0.12),
            )
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            _run(p, text, size=17, color=DARK)

    def two_columns(self, data: dict):
        self.header(data["title"])
        top = 1.62
        col_w = (CONTENT_W - 0.36) / 2
        col_h = 4.55

        for idx, (title_key, list_key) in enumerate([("left_title", "left"), ("right_title", "right")]):
            x = ML + idx * (col_w + 0.36)
            card = _rect(self.slide, _in(x), _in(top), _in(col_w), _in(col_h), CARD, radius=True, line=BORDER)

            # coloured header strip on card
            _rect(self.slide, _in(x), _in(top), _in(col_w), _in(0.52), PRIMARY_LIGHT if idx == 0 else TEAL_LIGHT, radius=True)
            _rect(self.slide, _in(x), _in(top + 0.38), _in(col_w), _in(0.14), PRIMARY_LIGHT if idx == 0 else TEAL_LIGHT)

            hb = _text(self.slide, _in(x + 0.22), _in(top + 0.1), _in(col_w - 0.44), _in(0.38))
            _run(hb.text_frame.paragraphs[0], data[title_key], size=16, bold=True, color=PRIMARY if idx == 0 else TEAL)

            items = data[list_key]
            for j, item in enumerate(items):
                row_y = top + 0.68 + j * 0.72
                dot = _rect(
                    self.slide,
                    _in(x + 0.22),
                    _in(row_y + 0.12),
                    _in(0.14),
                    _in(0.14),
                    PRIMARY if idx == 0 else TEAL,
                    radius=True,
                )
                tb = _text(self.slide, _in(x + 0.48), _in(row_y), _in(col_w - 0.62), _in(0.62))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                _run(tf.paragraphs[0], item, size=15, color=DARK)

    def agenda(self, data: dict):
        self.canvas()
        self.header(data["title"])
        items = data["items"]
        card_w = (CONTENT_W - 0.48) / 2
        card_h = 0.92
        positions = [
            (ML, 1.72),
            (ML + card_w + 0.48, 1.72),
            (ML, 1.72 + card_h + 0.28),
            (ML + card_w + 0.48, 1.72 + card_h + 0.28),
            (ML + (CONTENT_W - card_w) / 2, 1.72 + 2 * (card_h + 0.28)),
        ]
        for (num, label, color), (x, y) in zip(items, positions):
            _rect(self.slide, _in(x), _in(y), _in(card_w), _in(card_h), CARD, radius=True, line=BORDER)
            nb = _rect(self.slide, _in(x + 0.2), _in(y + 0.22), _in(0.48), _in(0.48), color, radius=True)
            ntf = nb.text_frame
            ntf.clear()
            ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
            np = ntf.paragraphs[0]
            np.alignment = PP_ALIGN.CENTER
            _run(np, num, size=13, bold=True, color=WHITE)
            tb = _text(self.slide, _in(x + 0.82), _in(y + 0.18), _in(card_w - 1.0), _in(card_h - 0.36))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _run(tf.paragraphs[0], label, size=17, bold=True, color=DARK)
        self.footer()

    def cover(self, data: dict):
        _set_bg(self.slide, NAVY)
        _rect(self.slide, _in(0), _in(0), SLIDE_W, _in(0.1), TEAL)
        # decorative circles
        for cx, cy, r, alpha in [(11.2, 1.2, 1.8, PRIMARY_LIGHT), (1.5, 5.8, 1.2, TEAL), (10.5, 6.0, 0.7, PRIMARY)]:
            c = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, _in(cx), _in(cy), _in(r), _in(r))
            c.fill.solid()
            c.fill.fore_color.rgb = alpha
            c.line.fill.background()

        panel = _rect(self.slide, _in(1.5), _in(1.85), _in(10.33), _in(3.85), CARD, radius=True)
        # inner accent bar
        _rect(self.slide, _in(1.5), _in(1.85), SLIDE_W - _in(3.0), _in(0.08), PRIMARY, radius=True)

        tb = _text(self.slide, _in(2.0), _in(2.35), _in(9.33), _in(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        _run(p1, data["title"], size=42, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(10)
        _run(p2, data["subtitle"], size=20, color=PRIMARY, align=PP_ALIGN.CENTER)
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        p3.space_before = Pt(18)
        _run(p3, data["tagline"], size=15, color=SLATE, align=PP_ALIGN.CENTER)
        p4 = tf.add_paragraph()
        p4.alignment = PP_ALIGN.CENTER
        p4.space_before = Pt(24)
        _run(p4, data["footer"], size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    def closing(self, data: dict):
        _set_bg(self.slide, NAVY)
        _rect(self.slide, _in(0), _in(3.35), SLIDE_W, _in(0.06), TEAL)
        for cx, cy, r in [(1.0, 1.0, 1.5), (11.5, 5.5, 1.0)]:
            c = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, _in(cx), _in(cy), _in(r), _in(r))
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
            c.line.fill.background()

        tb = _text(self.slide, _in(1.2), _in(2.5), _in(10.93), _in(3.0))
        tf = tb.text_frame
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        _run(p1, data["title"], size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for line in data.get("bullets", []):
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(14)
            _run(p, line, size=18, color=SLATE_LIGHT, align=PP_ALIGN.CENTER)
        pe = tf.add_paragraph()
        pe.alignment = PP_ALIGN.CENTER
        pe.space_before = Pt(32)
        _run(pe, data.get("footer", ""), size=16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    def bullets_slide(self, data: dict):
        self.canvas()
        self.header(data["title"], data.get("intro"))
        self.bullet_rows(data["bullets"])
        self.footer()

    def module_slide(self, data: dict):
        self.canvas()
        self.header(
            data["title"],
            data.get("intro"),
            badge=data.get("badge"),
            badge_color=data.get("badge_color", PRIMARY),
        )
        self.bullet_rows(data["bullets"], top=1.62 if data.get("intro") else 1.52)
        self.footer()


def _safe_save_pptx(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(path)
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "-updated" + path.suffix)
        prs.save(alt)
        print(f"Note: {path.name} is open — saved to {alt.name} instead. Close PowerPoint and re-run.")
        return alt


def build_pptx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for i, data in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        b = SlideBuilder(slide, i, total)
        kind = data["kind"]

        if kind == "cover":
            b.cover(data)
        elif kind == "closing":
            b.closing(data)
        elif kind == "agenda":
            b.agenda(data)
        elif kind == "two_col":
            b.canvas()
            b.two_columns(data)
            b.footer()
        elif kind == "module":
            b.module_slide(data)
        else:
            b.bullets_slide(data)

    return _safe_save_pptx(prs, PPTX_PATH)


# ── PDF (matching structure) ──────────────────────────────────────────────────
def _pdf_styles(base):
    return {
        "CoverTitle": ParagraphStyle("CoverTitle", parent=base["Heading1"], fontSize=34, leading=40, textColor=colors.white, alignment=TA_CENTER),
        "CoverSub": ParagraphStyle("CoverSub", parent=base["Normal"], fontSize=18, leading=24, textColor=colors.HexColor("#93C5FD"), alignment=TA_CENTER),
        "CoverTag": ParagraphStyle("CoverTag", parent=base["Normal"], fontSize=13, leading=18, textColor=colors.HexColor("#CBD5E1"), alignment=TA_CENTER),
        "CoverFoot": ParagraphStyle("CoverFoot", parent=base["Normal"], fontSize=12, leading=16, textColor=PDF_TEAL, alignment=TA_CENTER),
        "Title": ParagraphStyle("Title", parent=base["Heading1"], fontSize=24, leading=30, textColor=PDF_NAVY, spaceAfter=4),
        "Intro": ParagraphStyle("Intro", parent=base["Normal"], fontSize=11.5, leading=15, textColor=PDF_SLATE),
        "Bullet": ParagraphStyle("Bullet", parent=base["Normal"], fontSize=12.5, leading=17, textColor=colors.HexColor("#1E293B"), leftIndent=18, spaceBefore=4, spaceAfter=4),
        "AgendaItem": ParagraphStyle("AgendaItem", parent=base["Normal"], fontSize=12, leading=16, textColor=colors.HexColor("#1E293B")),
        "ColHead": ParagraphStyle("ColHead", parent=base["Normal"], fontSize=12, leading=16, textColor=PDF_PRIMARY),
        "Footer": ParagraphStyle("Footer", parent=base["Normal"], fontSize=9, textColor=PDF_SLATE, alignment=TA_CENTER),
        "CloseTitle": ParagraphStyle("CloseTitle", parent=base["Heading1"], fontSize=26, leading=32, textColor=colors.white, alignment=TA_CENTER),
        "CloseLine": ParagraphStyle("CloseLine", parent=base["Normal"], fontSize=13, leading=18, textColor=colors.HexColor("#CBD5E1"), alignment=TA_CENTER),
        "CloseFoot": ParagraphStyle("CloseFoot", parent=base["Normal"], fontSize=13, textColor=PDF_TEAL, alignment=TA_CENTER),
    }


def _pdf_bullets(story, styles, bullets: list[str]) -> None:
    rows = [[Paragraph(f"<font color='#2563EB'>●</font>&nbsp;&nbsp;{b}", styles["Bullet"])] for b in bullets]
    t = Table(rows, colWidths=[7.8 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), PDF_CARD),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROUNDEDCORNERS", [6, 6, 6, 6]),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)


def build_pdf() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    page_size = landscape((10 * inch, 7.5 * inch))
    doc = SimpleDocTemplate(
        str(PDF_PATH), pagesize=page_size,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch,
        topMargin=0.5 * inch, bottomMargin=0.45 * inch,
        title="Jagiri's Kutumbam Overview", author="Jagiri Family",
    )
    styles = _pdf_styles(getSampleStyleSheet())
    story = []
    total = len(SLIDES)

    # Cover
    d0 = SLIDES[0]
    story += [Spacer(1, 1.35 * inch), Paragraph(d0["title"], styles["CoverTitle"]), Spacer(1, 0.12 * inch),
              Paragraph(d0["subtitle"], styles["CoverSub"]), Spacer(1, 0.2 * inch),
              Paragraph(d0["tagline"], styles["CoverTag"]), Spacer(1, 0.35 * inch),
              Paragraph(d0["footer"], styles["CoverFoot"]), PageBreak()]

    for i, data in enumerate(SLIDES[1:], start=2):
        kind = data["kind"]
        if kind == "closing":
            story += [Spacer(1, 1.5 * inch), Paragraph(data["title"], styles["CloseTitle"]), Spacer(1, 0.15 * inch)]
            for line in data.get("bullets", []):
                story.append(Paragraph(line, styles["CloseLine"]))
            story += [Spacer(1, 0.3 * inch), Paragraph(data.get("footer", ""), styles["CloseFoot"])]
            continue

        title = data["title"]
        if kind == "module" and data.get("badge"):
            title = f"[{data['badge']}]  {title}"
        story.append(Paragraph(title, styles["Title"]))
        if data.get("intro"):
            story += [Spacer(1, 0.06 * inch), Paragraph(data["intro"], styles["Intro"])]
        story.append(Spacer(1, 0.14 * inch))

        if kind == "agenda":
            rows = [[Paragraph(f"<b>{n}</b>", styles["AgendaItem"]), Paragraph(l, styles["AgendaItem"])] for n, l, _ in data["items"]]
            t = Table(rows, colWidths=[0.55 * inch, 6.8 * inch])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), PDF_CARD),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.append(t)
        elif kind == "two_col":
            max_len = max(len(data["left"]), len(data["right"]))
            left = data["left"] + [""] * (max_len - len(data["left"]))
            right = data["right"] + [""] * (max_len - len(data["right"]))
            rows = [[Paragraph(f"<b>{data['left_title']}</b>", styles["ColHead"]), Paragraph(f"<b>{data['right_title']}</b>", styles["ColHead"])]]
            for l, r in zip(left, right):
                rows.append([Paragraph(f"● {l}" if l else "", styles["Bullet"]), Paragraph(f"● {r}" if r else "", styles["Bullet"])])
            t = Table(rows, colWidths=[3.75 * inch, 3.75 * inch])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), PDF_LIGHT),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            story.append(t)
        else:
            _pdf_bullets(story, styles, data.get("bullets", []))

        story += [Spacer(1, 0.25 * inch), Paragraph(f"Jagiri's Kutumbam  ·  {i} / {total}", styles["Footer"]), PageBreak()]

    def draw_bg(canvas, doc_obj):
        canvas.saveState()
        pw, ph = page_size
        if doc_obj.page == 1 or doc_obj.page == total:
            canvas.setFillColor(PDF_NAVY)
            canvas.rect(0, 0, pw, ph, fill=1, stroke=0)
            if doc_obj.page == 1:
                canvas.setFillColor(PDF_TEAL)
                canvas.rect(0, ph - 12, pw, 12, fill=1, stroke=0)
        else:
            canvas.setFillColor(PDF_LIGHT)
            canvas.rect(0, 0, pw, ph, fill=1, stroke=0)
            canvas.setFillColor(PDF_PRIMARY)
            canvas.rect(0, ph - doc_obj.topMargin - 36, 5, 32, fill=1, stroke=0)
            canvas.setFillColor(PDF_PRIMARY)
            canvas.rect(0, ph - doc_obj.topMargin - 36, pw, 2, fill=1, stroke=0)
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_bg, onLaterPages=draw_bg)
    return PDF_PATH


def main() -> None:
    pptx = build_pptx()
    pdf = build_pdf()
    print(f"Created: {pptx}")
    print(f"Created: {pdf}")


if __name__ == "__main__":
    main()
