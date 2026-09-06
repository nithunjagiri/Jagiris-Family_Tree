"""Generate EN/TE Word (.docx) narration scripts for the 3-min walkthrough."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parent


def set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    # Better Telugu glyph fallback on Windows
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Nirmala UI")
    rfonts.set(qn("w:cs"), "Nirmala UI")


def add_p(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_cue_table(doc: Document, headers: tuple[str, str, str], rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val


def write_txt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


EN_TELE = """Welcome to Jagiri's Kutumbam — our family app. This is one place where our family can stay connected: see relatives, share photos, plan events, find places on the map, and message each other. Open the app, enter your username and password, and tap Login. If you forget your password, use Forgot password on this same screen.

After login, you land on the Dashboard. Here you get a quick family overview — member counts, upcoming birthdays, recent activity, and announcements. Use the menu on the side — or the menu icon on phone — to open every module. Think of the Dashboard as your home base.

Tap Family Members. You'll see the full family list. You can search and filter to find someone quickly. Open any person to view their profile — photo, details, and related information. If you have permission, you can also add a new member or edit an existing one from here. Always enter names and dates carefully so the family record stays accurate.

Next, open Family Tree. This shows how people are connected — parents, children, and relatives — in a visual tree, so it's easier to understand relationships than reading a long list.

Open Photo Gallery. Browse family albums and photos. Tap a picture to view it larger. To share a memory, use Upload, choose a photo, add a short note if you like, and save. Everyone with access can enjoy it later.

Open Events. Here you'll find family gatherings, celebrations, and important dates. Tap an event to see the details. If you're organizing something, create an event so the family knows the date, time, and what to expect.

Open Places and map. This shows where family members are connected on a map. Use it to understand locations at a glance — helpful when planning visits or events.

Need someone fast? Open Search, type a name, and jump straight to their page. For private conversations, open Messages. Pick a chat, type your message, and send. You can also share images in chat when you want to send a photo quickly.

Finally, open Account and privacy to manage your own settings and privacy preferences. If you need help from the family admins, use Contact us from the menu.

That's Jagiri's Kutumbam in three minutes: Dashboard to start, Members and Tree to know the family, Gallery and Events to share life, Places and Search to find people, and Messages to stay in touch. Thank you — open the app anytime, and keep our family connected."""

TE_TELE = """నమస్కారం. ఇది జగిరి'స్ కుటుంబం — మన ఫ్యామిలీ యాప్. ఇక్కడ ఒకే చోట మన కుటుంబం కలిసి ఉండవచ్చు: బంధువులను చూడటం, ఫోటోలు పంచుకోవడం, ఈవెంట్స్ ప్లాన్ చేయడం, మ్యాప్‌లో స్థలాలు చూడటం, మరియు మెసేజ్‌లు పంపడం. యాప్ తెరిచి, మీ username మరియు password ఇచ్చి Login నొక్కండి. Password మర్చిపోతే, అదే స్క్రీన్‌లో Forgot password వాడండి.

లాగిన్ అయిన తర్వాత Dashboard కనిపిస్తుంది. ఇక్కడ కుటుంబం గురించి త్వరిత సమాచారం ఉంటుంది — సభ్యుల సంఖ్య, రాబోయే పుట్టినరోజులు, ఇటీవలి కార్యకలాపాలు, మరియు ప్రకటనలు. సైడ్‌లో ఉన్న మెనూ — ఫోన్‌లో మెనూ ఐకాన్ — వాడి అన్ని విభాగాలు తెరవండి. Dashboard ను మీ హోమ్ బేస్‌గా భావించండి.

Family Members నొక్కండి. పూర్తి కుటుంబ జాబితా కనిపిస్తుంది. ఎవరినైనా త్వరగా కనుక్కోవడానికి search మరియు filter వాడవచ్చు. ఎవరి పేరు మీదైనా నొక్కి వారి profile చూడండి — ఫోటో, వివరాలు, సంబంధిత సమాచారం. మీకు అనుమతి ఉంటే, ఇక్కడ నుంచే కొత్త సభ్యుడిని add చేయవచ్చు లేదా ఉన్న వారిని edit చేయవచ్చు. పేర్లు, తేదీలు జాగ్రత్తగా నమోదు చేయండి — కుటుంబ రికార్డు సరిగ్గా ఉండాలి.

ఇప్పుడు Family Tree తెరవండి. ఇక్కడ తల్లిదండ్రులు, పిల్లలు, బంధువులు ఎలా అనుసంధానమై ఉన్నారో చెట్టు రూపంలో కనిపిస్తుంది. జాబితా చదవడం కంటే సంబంధాలు అర్థం చేసుకోవడం సులభం.

Photo Gallery తెరవండి. కుటుంబ ఆల్బమ్‌లు మరియు ఫోటోలు చూడండి. పెద్దగా చూడాలంటే ఫోటో మీద నొక్కండి. జ్ఞాపకం పంచుకోవాలంటే Upload వాడి ఫోటో ఎంచుకోండి, కావాలంటే చిన్న నోట్ రాసి save చేయండి. యాక్సెస్ ఉన్నవారంతా తర్వాత చూడవచ్చు.

Events తెరవండి. ఇక్కడ కుటుంబ కలయికలు, వేడుకలు, ముఖ్యమైన తేదీలు ఉంటాయి. వివరాలు చూడాలంటే ఈవెంట్ మీద నొక్కండి. మీరు ఏదైనా నిర్వహిస్తుంటే, ఈవెంట్ create చేయండి — తేదీ, సమయం, వివరాలు కుటుంబానికి తెలుస్తాయి.

Places & map తెరవండి. మ్యాప్‌పై కుటుంబ సభ్యుల స్థలాలు కనిపిస్తాయి. సందర్శనలు లేదా ఈవెంట్స్ ప్లాన్ చేసేటప్పుడు లొకేషన్‌లు త్వరగా అర్థమవుతాయి.

ఎవరినైనా త్వరగా కావాలా? Search తెరిచి పేరు టైప్ చేసి వారి పేజీకి వెళ్లండి. వ్యక్తిగత సంభాషణకు Messages తెరవండి. చాట్ ఎంచుకుని మెసేజ్ టైప్ చేసి పంపండి. ఫోటో పంపాలంటే చాట్‌లోనే images కూడా share చేయవచ్చు.

చివరగా, Account & privacy తెరిచి మీ సెట్టింగ్స్ మరియు ప్రైవసీ ఆప్షన్స్ నిర్వహించుకోండి. అడ్మిన్‌ల సహాయం కావాలంటే మెనూ నుంచి Contact us వాడండి.

మూడు నిమిషాల్లో జగిరి'స్ కుటుంబం ఇదే: మొదలు Dashboard, కుటుంబం తెలుసుకోవడానికి Members మరియు Tree, జీవితం పంచుకోవడానికి Gallery మరియు Events, వ్యక్తులను కనుక్కోవడానికి Places మరియు Search, సంబంధం ఉంచుకోవడానికి Messages. ధన్యవాదాలు — ఎప్పుడైనా యాప్ తెరిచి, మన కుటుంబాన్ని కలిపి ఉంచుదాం."""


def build_english_docx() -> Path:
    doc = Document()
    set_normal_style(doc)
    doc.add_heading("Jagiri's Kutumbam — 3-Minute App Walkthrough (English)", level=0)
    add_p(doc, "Purpose: Word-for-word narration for a short demo video.")
    add_p(doc, "Length: about 2 minutes 45 seconds to 3 minutes.")
    add_p(doc, "Audience: Family members using the web or Android app for the first time.")

    doc.add_heading("Before you record", level=1)
    add_bullets(
        doc,
        [
            "Log in with a demo (or safe) account — hide real private phone numbers if possible.",
            "Have sample data ready: a few members, one photo, one event, one chat thread.",
            "Use landscape (phone sideways) or a computer browser window.",
            "Speak slightly slower than normal; pause about 1 second after each screen change.",
        ],
    )

    doc.add_heading("Recording cue sheet", level=1)
    add_cue_table(
        doc,
        ("Time", "Screen to show", "Narration block"),
        [
            ("0:00", "App logo / Login", "Opening"),
            ("0:20", "Dashboard", "Dashboard"),
            ("0:45", "Family Members → one profile", "Members"),
            ("1:10", "Family Tree", "Tree"),
            ("1:25", "Photo Gallery", "Gallery"),
            ("1:40", "Events", "Events"),
            ("1:55", "Places & map", "Places"),
            ("2:10", "Search → Messages", "Search & chat"),
            ("2:30", "Account & privacy (quick)", "Account"),
            ("2:40", "Dashboard again / logo", "Closing"),
        ],
    )

    doc.add_heading("Full narration script (word-for-word)", level=1)
    sections = [
        (
            "Opening (about 20 sec)",
            "Welcome to Jagiri's Kutumbam — our family app.\n\n"
            "This is one place where our family can stay connected: see relatives, share photos, plan events, find places on the map, and message each other.\n\n"
            "Open the app, enter your username and password, and tap Login. If you forget your password, use Forgot password on this same screen.",
        ),
        (
            "Dashboard (about 25 sec)",
            "After login, you land on the Dashboard.\n\n"
            "Here you get a quick family overview — member counts, upcoming birthdays, recent activity, and announcements.\n\n"
            "Use the menu on the side — or the menu icon on phone — to open every module. Think of the Dashboard as your home base.",
        ),
        (
            "Family Members (about 25 sec)",
            "Tap Family Members.\n\n"
            "You'll see the full family list. You can search and filter to find someone quickly.\n\n"
            "Open any person to view their profile — photo, details, and related information.\n\n"
            "If you have permission, you can also add a new member or edit an existing one from here. Always enter names and dates carefully so the family record stays accurate.",
        ),
        (
            "Family Tree (about 15 sec)",
            "Next, open Family Tree.\n\n"
            "This shows how people are connected — parents, children, and relatives — in a visual tree, so it's easier to understand relationships than reading a long list.",
        ),
        (
            "Photo Gallery (about 15 sec)",
            "Open Photo Gallery.\n\n"
            "Browse family albums and photos. Tap a picture to view it larger.\n\n"
            "To share a memory, use Upload, choose a photo, add a short note if you like, and save. Everyone with access can enjoy it later.",
        ),
        (
            "Events (about 15 sec)",
            "Open Events.\n\n"
            "Here you'll find family gatherings, celebrations, and important dates. Tap an event to see the details.\n\n"
            "If you're organizing something, create an event so the family knows the date, time, and what to expect.",
        ),
        (
            "Places & map (about 15 sec)",
            "Open Places & map.\n\n"
            "This shows where family members are connected on a map. Use it to understand locations at a glance — helpful when planning visits or events.",
        ),
        (
            "Search & Messages (about 20 sec)",
            "Need someone fast? Open Search, type a name, and jump straight to their page.\n\n"
            "For private conversations, open Messages. Pick a chat, type your message, and send. You can also share images in chat when you want to send a photo quickly.",
        ),
        (
            "Account (about 10 sec)",
            "Finally, open Account & privacy to manage your own settings and privacy preferences.\n\n"
            "If you need help from the family admins, use Contact us from the menu.",
        ),
        (
            "Closing (about 15 sec)",
            "That's Jagiri's Kutumbam in three minutes:\n\n"
            "Dashboard to start, Members and Tree to know the family, Gallery and Events to share life, Places and Search to find people, and Messages to stay in touch.\n\n"
            "Thank you — open the app anytime, and keep our family connected.",
        ),
    ]
    for title, body in sections:
        doc.add_heading(title, level=2)
        for para in body.split("\n\n"):
            add_p(doc, para)

    doc.add_heading("Teleprompter copy (no headings)", level=1)
    for para in EN_TELE.split("\n\n"):
        add_p(doc, para)

    doc.add_heading("How to create the video", level=1)
    doc.add_heading("Option A — Phone", level=2)
    add_bullets(
        doc,
        [
            "Install the Android app (or open the website in Chrome).",
            "Turn on Screen recording (notification shade → Screen record). Enable microphone.",
            "Put the teleprompter text on a second device (or print it).",
            "Record while tapping through each screen in cue sheet order.",
            "Stop recording. Trim in Google Photos Edit or CapCut.",
            "Add a title card and end card with the app/website link.",
            "Export and share (WhatsApp, Drive, YouTube Unlisted).",
        ],
    )
    doc.add_heading("Option B — Windows PC", level=2)
    add_bullets(
        doc,
        [
            "Open the app in Chrome (full screen: F11).",
            "Press Win + G → start Capture (include microphone).",
            "Narrate using the teleprompter copy while clicking each menu item.",
            "Edit in Clipchamp or CapCut desktop.",
            "Export MP4 (1080p if possible).",
        ],
    )
    doc.add_heading("Option C — Voice-over after recording", level=2)
    add_bullets(
        doc,
        [
            "Record the screen without talking (silent demo).",
            "Record voice separately on your phone using the teleprompter text.",
            "In CapCut/Clipchamp: import video + audio, align, export.",
        ],
    )

    path = OUT / "Jagiris-Kutumbam-3min-narration-EN.docx"
    doc.save(path)
    return path


def build_telugu_docx() -> Path:
    doc = Document()
    set_normal_style(doc)
    doc.add_heading("జగిరి'స్ కుటుంబం — 3 నిమిషాల యాప్ వాక్‌త్రూ (తెలుగు)", level=0)
    add_p(doc, "ఉద్దేశం: షార్ట్ డెమో వీడియో కోసం పూర్తి న్యారేషన్ స్క్రిప్ట్.")
    add_p(doc, "నిడివి: సుమారు 2 నిమిషాల 45 సెకన్ల నుంచి 3 నిమిషాలు.")
    add_p(doc, "ప్రేక్షకులు: మొదటిసారి వెబ్ లేదా Android యాప్ వాడే కుటుంబ సభ్యులు.")
    add_p(
        doc,
        "గమనిక: యాప్‌లో కనిపించే మెనూ పేర్లు (Login, Dashboard, Family Members మొదలైనవి) ఇంగ్లీష్‌లోనే ఉంచాం — స్క్రీన్‌పై అలాగే కనిపిస్తాయి.",
    )

    doc.add_heading("రికార్డ్ చేసే ముందు", level=1)
    add_bullets(
        doc,
        [
            "డెమో లేదా సేఫ్ అకౌంట్‌తో లాగిన్ అవ్వండి — అవసరమైతే ఫోన్ నంబర్లు దాచండి.",
            "కొంత సాంపిల్ డేటా సిద్ధంగా ఉంచండి: కొందరు సభ్యులు, ఒక ఫోటో, ఒక ఈవెంట్, ఒక చాట్.",
            "ఫోన్‌ను ల్యాండ్‌స్కేప్‌లో లేదా కంప్యూటర్ బ్రౌజర్‌లో చూపండి.",
            "కాస్త నెమ్మదిగా మాట్లాడండి; ప్రతి స్క్రీన్ మారిన తర్వాత సుమారు 1 సెకను ఆగండి.",
        ],
    )

    doc.add_heading("రికార్డింగ్ క్యూ షీట్", level=1)
    add_cue_table(
        doc,
        ("సమయం", "చూపించాల్సిన స్క్రీన్", "భాగం"),
        [
            ("0:00", "యాప్ లోగో / Login", "ప్రారంభం"),
            ("0:20", "Dashboard", "డాష్‌బోర్డ్"),
            ("0:45", "Family Members → ప్రొఫైల్", "సభ్యులు"),
            ("1:10", "Family Tree", "వంశవృక్షం"),
            ("1:25", "Photo Gallery", "గ్యాలరీ"),
            ("1:40", "Events", "ఈవెంట్స్"),
            ("1:55", "Places & map", "ప్లేసెస్"),
            ("2:10", "Search → Messages", "సెర్చ్ & చాట్"),
            ("2:30", "Account & privacy", "అకౌంట్"),
            ("2:40", "Dashboard / లోగో", "ముగింపు"),
        ],
    )

    doc.add_heading("పూర్తి న్యారేషన్ స్క్రిప్ట్ (మాట్లాడేలా)", level=1)
    sections = [
        (
            "ప్రారంభం (సుమారు 20 సెకన్లు)",
            "నమస్కారం. ఇది జగిరి'స్ కుటుంబం — మన ఫ్యామిలీ యాప్.\n\n"
            "ఇక్కడ ఒకే చోట మన కుటుంబం కలిసి ఉండవచ్చు: బంధువులను చూడటం, ఫోటోలు పంచుకోవడం, ఈవెంట్స్ ప్లాన్ చేయడం, మ్యాప్‌లో స్థలాలు చూడటం, మరియు మెసేజ్‌లు పంపడం.\n\n"
            "యాప్ తెరిచి, మీ username మరియు password ఇచ్చి Login నొక్కండి. Password మర్చిపోతే, అదే స్క్రీన్‌లో Forgot password వాడండి.",
        ),
        (
            "డాష్‌బోర్డ్ (సుమారు 25 సెకన్లు)",
            "లాగిన్ అయిన తర్వాత Dashboard కనిపిస్తుంది.\n\n"
            "ఇక్కడ కుటుంబం గురించి త్వరిత సమాచారం ఉంటుంది — సభ్యుల సంఖ్య, రాబోయే పుట్టినరోజులు, ఇటీవలి కార్యకలాపాలు, మరియు ప్రకటనలు.\n\n"
            "సైడ్‌లో ఉన్న మెనూ — ఫోన్‌లో మెనూ ఐకాన్ — వాడి అన్ని విభాగాలు తెరవండి. Dashboard ను మీ హోమ్ బేస్‌గా భావించండి.",
        ),
        (
            "Family Members (సుమారు 25 సెకన్లు)",
            "Family Members నొక్కండి.\n\n"
            "పూర్తి కుటుంబ జాబితా కనిపిస్తుంది. ఎవరినైనా త్వరగా కనుక్కోవడానికి search మరియు filter వాడవచ్చు.\n\n"
            "ఎవరి పేరు మీదైనా నొక్కి వారి profile చూడండి — ఫోటో, వివరాలు, సంబంధిత సమాచారం.\n\n"
            "మీకు అనుమతి ఉంటే, ఇక్కడ నుంచే కొత్త సభ్యుడిని add చేయవచ్చు లేదా ఉన్న వారిని edit చేయవచ్చు. పేర్లు, తేదీలు జాగ్రత్తగా నమోదు చేయండి — కుటుంబ రికార్డు సరిగ్గా ఉండాలి.",
        ),
        (
            "Family Tree (సుమారు 15 సెకన్లు)",
            "ఇప్పుడు Family Tree తెరవండి.\n\n"
            "ఇక్కడ తల్లిదండ్రులు, పిల్లలు, బంధువులు ఎలా అనుసంధానమై ఉన్నారో చెట్టు రూపంలో కనిపిస్తుంది. జాబితా చదవడం కంటే సంబంధాలు అర్థం చేసుకోవడం సులభం.",
        ),
        (
            "Photo Gallery (సుమారు 15 సెకన్లు)",
            "Photo Gallery తెరవండి.\n\n"
            "కుటుంబ ఆల్బమ్‌లు మరియు ఫోటోలు చూడండి. పెద్దగా చూడాలంటే ఫోటో మీద నొక్కండి.\n\n"
            "జ్ఞాపకం పంచుకోవాలంటే Upload వాడి ఫోటో ఎంచుకోండి, కావాలంటే చిన్న నోట్ రాసి save చేయండి. యాక్సెస్ ఉన్నవారంతా తర్వాత చూడవచ్చు.",
        ),
        (
            "Events (సుమారు 15 సెకన్లు)",
            "Events తెరవండి.\n\n"
            "ఇక్కడ కుటుంబ కలయికలు, వేడుకలు, ముఖ్యమైన తేదీలు ఉంటాయి. వివరాలు చూడాలంటే ఈవెంట్ మీద నొక్కండి.\n\n"
            "మీరు ఏదైనా నిర్వహిస్తుంటే, ఈవెంట్ create చేయండి — తేదీ, సమయం, వివరాలు కుటుంబానికి తెలుస్తాయి.",
        ),
        (
            "Places & map (సుమారు 15 సెకన్లు)",
            "Places & map తెరవండి.\n\n"
            "మ్యాప్‌పై కుటుంబ సభ్యుల స్థలాలు కనిపిస్తాయి. సందర్శనలు లేదా ఈవెంట్స్ ప్లాన్ చేసేటప్పుడు లొకేషన్‌లు త్వరగా అర్థమవుతాయి.",
        ),
        (
            "Search & Messages (సుమారు 20 సెకన్లు)",
            "ఎవరినైనా త్వరగా కావాలా? Search తెరిచి పేరు టైప్ చేసి వారి పేజీకి వెళ్లండి.\n\n"
            "వ్యక్తిగత సంభాషణకు Messages తెరవండి. చాట్ ఎంచుకుని మెసేజ్ టైప్ చేసి పంపండి. ఫోటో పంపాలంటే చాట్‌లోనే images కూడా share చేయవచ్చు.",
        ),
        (
            "Account (సుమారు 10 సెకన్లు)",
            "చివరగా, Account & privacy తెరిచి మీ సెట్టింగ్స్ మరియు ప్రైవసీ ఆప్షన్స్ నిర్వహించుకోండి.\n\n"
            "అడ్మిన్‌ల సహాయం కావాలంటే మెనూ నుంచి Contact us వాడండి.",
        ),
        (
            "ముగింపు (సుమారు 15 సెకన్లు)",
            "మూడు నిమిషాల్లో జగిరి'స్ కుటుంబం ఇదే:\n\n"
            "మొదలు Dashboard, కుటుంబం తెలుసుకోవడానికి Members మరియు Tree, జీవితం పంచుకోవడానికి Gallery మరియు Events, వ్యక్తులను కనుక్కోవడానికి Places మరియు Search, సంబంధం ఉంచుకోవడానికి Messages.\n\n"
            "ధన్యవాదాలు — ఎప్పుడైనా యాప్ తెరిచి, మన కుటుంబాన్ని కలిపి ఉంచుదాం.",
        ),
    ]
    for title, body in sections:
        doc.add_heading(title, level=2)
        for para in body.split("\n\n"):
            add_p(doc, para)

    doc.add_heading("టెలిప్రాంప్టర్ కాపీ (హెడ్డింగ్‌లు లేకుండా)", level=1)
    for para in TE_TELE.split("\n\n"):
        add_p(doc, para)

    doc.add_heading("వీడియో ఎలా తయారు చేయాలి", level=1)
    doc.add_heading("ఆప్షన్ A — ఫోన్", level=2)
    add_bullets(
        doc,
        [
            "Android యాప్ ఇన్‌స్టాల్ చేయండి (లేదా Chromeలో వెబ్‌సైట్ తెరవండి).",
            "Screen recording ఆన్ చేయండి; మైక్రోఫోన్ ఆన్ చేయండి.",
            "టెలిప్రాంప్టర్ టెక్స్ట్ మరో ఫోన్‌లో లేదా ప్రింట్‌లో ఉంచుకోండి.",
            "క్యూ షీట్ క్రమంలో స్క్రీన్లు తెరిచి రికార్డ్ చేయండి.",
            "ఆపి, Google Photos Edit లేదా CapCutలో కత్తిరించండి.",
            "టైటిల్ కార్డ్ మరియు ఎండ్ కార్డ్ జోడించండి.",
            "Export చేసి WhatsApp / Drive / YouTube Unlistedలో షేర్ చేయండి.",
        ],
    )
    doc.add_heading("ఆప్షన్ B — Windows PC", level=2)
    add_bullets(
        doc,
        [
            "Chromeలో యాప్ తెరవండి (F11 ఫుల్ స్క్రీన్).",
            "Win + G → Capture ప్రారంభించండి (మైక్‌తో).",
            "టెలిప్రాంప్టర్ చదువుతూ మెనూ క్లిక్ చేయండి.",
            "Clipchamp లేదా CapCutలో ఎడిట్ చేయండి.",
            "MP4గా export చేయండి (సాధ్యమైతే 1080p).",
        ],
    )
    doc.add_heading("ఆప్షన్ C — ముందు వీడియో, తర్వాత వాయిస్", level=2)
    add_bullets(
        doc,
        [
            "మాట్లాడకుండా స్క్రీన్ రికార్డ్ చేయండి.",
            "టెలిప్రాంప్టర్‌తో వేరుగా వాయిస్ రికార్డ్ చేయండి.",
            "CapCut/Clipchampలో వీడియో + ఆడియో కలిపి export చేయండి.",
        ],
    )

    path = OUT / "Jagiris-Kutumbam-3min-narration-TE.docx"
    doc.save(path)
    return path


def build_txt_files() -> tuple[Path, Path]:
    en_txt = OUT / "Jagiris-Kutumbam-3min-narration-EN.txt"
    te_txt = OUT / "Jagiris-Kutumbam-3min-narration-TE.txt"

    write_txt(
        en_txt,
        "Jagiri's Kutumbam — 3-Minute App Walkthrough (English)\n"
        "=====================================================\n\n"
        "TELEPROMPTER SCRIPT\n"
        "-------------------\n\n"
        f"{EN_TELE}\n\n"
        "CUE SHEET\n"
        "---------\n"
        "0:00  Login\n"
        "0:20  Dashboard\n"
        "0:45  Family Members → profile\n"
        "1:10  Family Tree\n"
        "1:25  Photo Gallery\n"
        "1:40  Events\n"
        "1:55  Places & map\n"
        "2:10  Search → Messages\n"
        "2:30  Account & privacy\n"
        "2:40  Closing\n",
    )

    write_txt(
        te_txt,
        "జగిరి'స్ కుటుంబం — 3 నిమిషాల యాప్ వాక్‌త్రూ (తెలుగు)\n"
        "=====================================================\n\n"
        "గమనిక: యాప్ మెనూ పేర్లు (Login, Dashboard, Family Members మొదలైనవి) ఇంగ్లీష్‌లోనే — స్క్రీన్‌పై అలాగే కనిపిస్తాయి.\n\n"
        "టెలిప్రాంప్టర్ స్క్రిప్ట్\n"
        "-------------------\n\n"
        f"{TE_TELE}\n\n"
        "క్యూ షీట్\n"
        "---------\n"
        "0:00  Login\n"
        "0:20  Dashboard\n"
        "0:45  Family Members → ప్రొఫైల్\n"
        "1:10  Family Tree\n"
        "1:25  Photo Gallery\n"
        "1:40  Events\n"
        "1:55  Places & map\n"
        "2:10  Search → Messages\n"
        "2:30  Account & privacy\n"
        "2:40  ముగింపు\n",
    )
    return en_txt, te_txt


def main() -> None:
    en_docx = build_english_docx()
    te_docx = build_telugu_docx()
    en_txt, te_txt = build_txt_files()
    print("Wrote:", en_docx)
    print("Wrote:", te_docx)
    print("Wrote:", en_txt)
    print("Wrote:", te_txt)


if __name__ == "__main__":
    main()
